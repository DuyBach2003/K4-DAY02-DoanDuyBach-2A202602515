"""Tổng hợp results.xlsx từ log của các lần chạy thật (không nhập số tay).

Chạy từ thư mục bài nộp, sau khi đã chạy eval.py score/grade vào eval_out/:
    python code/make_results.py
Nguồn số liệu:
  logs/<exp>/seed<k>/{config.json, summary.json}   (Bước 1, 2, 4 - chép từ runs/ bởi experiments.py)
  logs/step3_inference.json, logs/step3_latency.json (Bước 3)
  logs/step4_<exp>.json, eval_out/<tag>_*.{json,csv}, eval_out/grade_I.json (Bước 4, eval.py)
  predictions/<exp>_seed<k>_test.csv (recall từng seed), logs/diag_precise_bn.json, logs/diag_factorial_val.json
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))

import train  # noqa: E402,F401  (đặt sys.path tới eval.py gốc)
from eval import compute_metrics, read_pred  # noqa: E402
from experiments import REGISTRY, TRAINING_META, FINAL_META  # noqa: E402

CLASS = ["Chinee Apple", "Lantana", "Parkinsonia", "Parthenium", "Prickly Acacia", "Rubber Vine", "Siam Weed",
         "Snake Weed", "Negatives"]
# cấu hình đo độ trễ (logs/step3_latency.json) ứng với phương pháp suy luận của từng cấu hình ở Bước 4
FINAL_LATENCY = {"T00": "I00_fp32_b1", "F01": "I04_res288_b1"}
# ghi chú định tính cho các dòng của Summary (số liệu nằm ở các cột và các sheet chi tiết)
SUMMARY_NOTES = {
    "I05": "không chọn: ngang thành viên tốt nhất (B02) mà độ trễ ~5× I00",
    "B02": "backbone tốt nhất ở 224 nhưng train chậm gần 2× DeiT-S; không đi tiếp (ngân sách GPU)",
    "I02c": "không chọn: không phân biệt được với I04_288 mà độ trễ ~2,7×",
    "I03c": "như I02c (gộp logit)",
    "I04_320": "không phân biệt được với I04_288 (bão hoà)",
    "I04_288": "CHỌN cho chung kết: 1 view, độ trễ ~2× I00, trong ngân sách thời gian thực",
    "T05": "Δ so với T00 chưa vượt nhiễu seed (xem sheet Training)",
    "T09": "công thức của chung kết F01 (kết hợp T02 + T05)",
}


def noise_std() -> float:
    """std mẫu (ddof=1) của macro-F1 val giữa các seed của T00 (cùng cấu hình, chỉ khác seed)."""
    v = [r[0]["val_macro_f1"] for s in (0, 1, 2) if (r := load_run("T00", s))]
    return float(np.std(v, ddof=1)) if len(v) >= 2 else np.nan


def noise_verdict(delta: float, sd: float) -> str:
    """So Δ của hai lần chạy 1 seed với std của hiệu hai lần chạy (≈ √2·sd)."""
    r = abs(delta) / (np.sqrt(2) * sd)
    if r <= 1:
        return "không phân biệt được (|Δ| ≤ √2·σ)"
    word = "tốt hơn" if delta > 0 else "kém hơn"
    return f"{word} rõ (|Δ| > 2√2·σ)" if r > 2 else f"có thể {word}, chưa chắc (1 seed)"


def load_run(exp: str, seed: int = 0):
    d = Path("logs") / exp / f"seed{seed}"
    if not (d / "summary.json").exists():
        return None
    s = json.loads((d / "summary.json").read_text())
    c = json.loads((d / "config.json").read_text())
    return s, c


def curve_name(exp: str) -> str:
    desc = REGISTRY.get(exp, {}).get("desc", "")
    return f"curves/{exp}_{desc}.png" if desc else f"curves/{exp}.png"


def sheet_backbones():
    bt = Path("logs/bench_train.json")
    bench = json.loads(bt.read_text())["results"] if bt.exists() else {}
    pb = Path("logs/diag_precise_bn.json")
    pbn = json.loads(pb.read_text()) if pb.exists() else {}
    rows = []
    for exp in sorted(e for e in REGISTRY if e.startswith(("B", "D"))):
        r = load_run(exp)
        if r is None:
            continue
        s, c = r
        lat = s.get("latency_b1_fp32") or {}
        rows.append({
            "exp_id": exp, "backbone": c["config"]["backbone"], "tag trọng số": s["weight_tag"],
            "#tham số (M)": round(s["params_m"], 2), "GMAC": round(s["gmacs"], 2),
            "độ phân giải": c["config"]["img_size"], "epoch": c["config"]["epochs"], "seed": s["seed"],
            "dtype train": c["amp_dtype"], "best epoch": s["best_epoch"],
            "macro-F1 val": s["val_macro_f1"], "top-1 val": s["val_top1"], "balanced acc val": s["val_bal_acc"],
            "F1 Chinee Apple val": s["val_f1_per_class"][0], "F1 Snake Weed val": s["val_f1_per_class"][7],
            "thời gian train/epoch (s)": round(s["train_time_per_epoch_s"], 1),
            "tốc độ train đo riêng (ảnh/s)": round(bench.get(exp, {}).get("train_img_per_s", np.nan), 1),
            "độ trễ batch-1 p50 (ms)": round(lat.get("p50", np.nan), 2),
            "độ trễ batch-1 p95 (ms)": round(lat.get("p95", np.nan), 2),
            # chẩn đoán precise-BN (code/diag_precise_bn.py): đánh giá lại FP32 trước/sau khi tính lại thống kê BN
            "precise-BN: macro-F1 val trước (FP32)": pbn.get(exp, {}).get("before", {}).get("macro_f1", np.nan),
            "precise-BN: macro-F1 val sau": pbn.get(exp, {}).get("after_precise_bn", {}).get("macro_f1", np.nan),
            "ảnh biểu đồ": curve_name(exp), "ghi chú": REGISTRY[exp].get("note", ""),
        })
    return pd.DataFrame(rows)


def sheet_training():
    rows = []
    base = load_run("T00")
    base_f1 = base[0]["val_macro_f1"] if base else np.nan
    base_top1 = base[0]["val_top1"] if base else np.nan
    sd = noise_std()
    # T00 seed 1, 2: cùng cấu hình, chỉ khác seed -> Δ của chúng chính là nhiễu seed
    runs = [(e, 0, m) for e, m in TRAINING_META.items()]
    runs[1:1] = [("T00", s, {"axis": "lặp T00 (đo nhiễu seed)", "diff": f"công thức nền, seed {s} (không khác gì)"})
                 for s in (1, 2)]
    for exp, seed, meta in runs:
        r = load_run(exp, seed)
        if r is None:
            continue
        s, c = r
        d = s["val_macro_f1"] - base_f1
        if exp == "T00":
            verdict = "mốc" if seed == 0 else "nhiễu seed (cùng cấu hình)"
        else:
            verdict = noise_verdict(d, sd)
        rows.append({
            "exp_id": exp, "backbone": c["config"]["backbone"], "trục": meta["axis"],
            "khác T00 ở điểm nào": meta["diff"], "seed": s["seed"], "best epoch": s["best_epoch"],
            "macro-F1 val": s["val_macro_f1"], "top-1 val": s["val_top1"],
            "Δ macro-F1 so với T00": d, "Δ top-1 so với T00": s["val_top1"] - base_top1,
            "Δ / σ seed": d / sd, "so với nhiễu": verdict,
            "F1 Chinee Apple val": s["val_f1_per_class"][0], "F1 Snake Weed val": s["val_f1_per_class"][7],
            "ECE val": s["val_ece"], "thời gian train/epoch (s)": round(s["train_time_per_epoch_s"], 1),
            "ảnh biểu đồ": curve_name(exp), "ghi chú": meta.get("note", ""),
        })
    return pd.DataFrame(rows)


def sheet_inference():
    p = Path("logs/step3_inference.json")
    if not p.exists():
        return pd.DataFrame(), pd.DataFrame()
    inf = json.loads(p.read_text())
    lat = json.loads(Path("logs/step3_latency.json").read_text())
    lat_key = {"I00": "I00_fp32", "I01": "I01_hflip_K2", "I03a": "I01_hflip_K2", "I02b": "I02b_10crop",
               "I03b": "I02b_10crop", "I02c": "I02c_3scale", "I03c": "I02c_3scale", "I04_224": "I00_fp32",
               "I04_256": "I04_res256", "I04_288": "I04_res288", "I04_320": "I04_res320", "I05": "I05_ensemble",
               "I06": "I00_fp32", "I07_I00": "I00_fp32", "I07_I01": "I01_hflip_K2", "I08a": "I08a_fusedBN_fp32",
               "I08b": "I08b_fp16", "I08c": "I08c_amp", "I09": "I00_fp32", "I09_I01": "I01_hflip_K2",
               "I09_I01_I07": "I01_hflip_K2", "I02a": None}
    base = lat.get("I00_fp32_b1", {}).get("p50", np.nan)
    rows = []
    for code, r in inf["results"].items():
        k = lat_key.get(code)
        l1 = lat.get(f"{k}_b1", {}) if k else {}
        l32 = lat.get(f"{k}_b32", {}) if k else {}
        rows.append({
            "exp_id": code, "phương pháp": r["method"], "mô hình/checkpoint": r.get("model", f"{inf['exp']} seed{inf['seed']}"),
            "K (view/mô hình)": r["k"], "macro-F1 val": r["macro_f1"], "top-1 val": r["top1"], "ECE val": r["ece"],
            "NLL val": r["nll"], "F1 Chinee Apple val": r["f1_chinee"], "F1 Snake Weed val": r["f1_snake"],
            "T": r.get("T", np.nan), "ECE val (T khớp chéo 2 nửa val)": r.get("ece_after_crossfit", np.nan),
            "độ trễ p50 b1 (ms)": l1.get("p50", np.nan), "độ trễ p95 b1 (ms)": l1.get("p95", np.nan),
            "độ trễ p99 b1 (ms)": l1.get("p99", np.nan), "thông lượng b32 (ảnh/s)": l32.get("images_per_s", np.nan),
            "chi phí tương đối so với I00": l1.get("p50", np.nan) / base if l1 else np.nan,
            "cấu hình đo độ trễ": k or "(không đo riêng)",
        })
    lat_rows = [{"cấu hình": v["config"], "thiết bị": v["gpu"], "dtype": v["dtype"], "batch": v["batch"],
                 "độ phân giải": v["img_size"], "gộp BN": "có" if v.get("fused_bn") else "không",
                 "p50 (ms)": v["p50"], "p95 (ms)": v["p95"], "p99 (ms)": v["p99"], "ảnh/s": v["images_per_s"],
                 "số lần đo": v.get("n", np.nan), "warmup": v.get("warmup", np.nan), "torch": v["torch"],
                 "tính tiền xử lý": "không", "ghi chú": v.get("note", "")} for v in lat.values()]
    return pd.DataFrame(rows), pd.DataFrame(lat_rows)


def test_metrics(path: Path) -> dict | None:
    """Chỉ số của một file dự đoán test (đọc bằng eval.read_pred, tính bằng eval.compute_metrics)."""
    if not path.exists():
        return None
    p = read_pred(str(path))
    return compute_metrics(p.y_true, p.y_pred, p.probs)


def eval_summary(tag: str) -> dict | None:
    p = Path(f"eval_out/{tag}_summary.json")
    return json.loads(p.read_text()) if p.exists() else None


def fmt_ms(e: dict, key: str, idx: int | None = None) -> str:
    m, s = (e[key]["mean"], e[key]["std"]) if idx is None else (e[key]["mean"][idx], e[key]["std"][idx])
    return f"{m:.4f} ± {s:.4f}"


def sheet_final():
    lat = json.loads(Path("logs/step3_latency.json").read_text()) if Path("logs/step3_latency.json").exists() else {}
    rows = []
    for exp, meta in FINAL_META.items():
        p = Path(f"logs/step4_{exp}.json")
        if not p.exists():
            continue
        st = json.loads(p.read_text())
        recs = [st[k] for k in sorted(st, key=int)]
        p95 = lat.get(FINAL_LATENCY.get(exp, ""), {}).get("p95", np.nan)
        for r in recs:
            m = test_metrics(Path(f"predictions/{exp}_seed{r['seed']}_test.csv")) or {}
            mu = test_metrics(Path(f"predictions/{exp}uncal_seed{r['seed']}_test.csv")) or {}
            rows.append({"exp_id": exp, "cấu hình": meta, "seed": r["seed"], "T (khớp trên val)": r["temperature"],
                         "macro-F1 val": r["val_macro_f1"], "macro-F1 test": r["test_macro_f1"],
                         "top-1 test": r["test_top1"], "ECE test": r["test_ece"],
                         "ECE test chưa TS (T = 1)": mu.get("ece", r["test_ece"] if r["temperature"] == 1 else np.nan),
                         "recall Chinee Apple test": m["recall"][0] if m else np.nan,
                         "recall Snake Weed test": m["recall"][7] if m else np.nan,
                         "F1 Chinee Apple test": m["f1"][0] if m else np.nan,
                         "F1 Snake Weed test": m["f1"][7] if m else np.nan,
                         "độ trễ p95 b1 (ms)": p95})
        agg = {"exp_id": f"{exp} (mean ± std, {len(recs)} seed)", "cấu hình": meta, "seed": "tổng hợp"}
        vf = [r["val_macro_f1"] for r in recs]
        agg["macro-F1 val"] = f"{np.mean(vf):.4f} ± {np.std(vf, ddof=1):.4f}"
        e = eval_summary(exp)
        if e:
            for col, key in (("macro-F1 test", "macro_f1"), ("top-1 test", "top1"), ("ECE test", "ece")):
                agg[col] = fmt_ms(e, key)
            eu = eval_summary(f"{exp}uncal")
            agg["ECE test chưa TS (T = 1)"] = fmt_ms(eu, "ece") if eu else agg["ECE test"]
            agg["recall Chinee Apple test"] = fmt_ms(e, "recall", 0)
            agg["recall Snake Weed test"] = fmt_ms(e, "recall", 7)
            agg["F1 Chinee Apple test"] = fmt_ms(e, "f1", 0)
            agg["F1 Snake Weed test"] = fmt_ms(e, "f1", 7)
        agg["độ trễ p95 b1 (ms)"] = p95
        rows.append(agg)
    return pd.DataFrame(rows)


def sheet_compare():
    """Khối đầu sheet Summary: chung kết so với mốc trên TEST (mean ± std qua seed, theo eval.py)."""
    lat = json.loads(Path("logs/step3_latency.json").read_text()) if Path("logs/step3_latency.json").exists() else {}
    base = eval_summary("T00")
    rows = []
    for exp, label in (("F01", "chung kết"), ("T00", "mốc (T00 + I00)")):
        e = eval_summary(exp)
        if not e:
            continue
        st = json.loads(Path(f"logs/step4_{exp}.json").read_text())
        vf = [v["val_macro_f1"] for v in st.values()]
        row = {"cấu hình": f"{exp} - {label}", "mô tả": FINAL_META[exp], "số seed": len(e["seeds"]),
               "macro-F1 val": f"{np.mean(vf):.4f} ± {np.std(vf, ddof=1):.4f}",
               "macro-F1 test": fmt_ms(e, "macro_f1"), "top-1 test": fmt_ms(e, "top1"), "ECE test": fmt_ms(e, "ece"),
               "recall Chinee Apple test": fmt_ms(e, "recall", 0), "recall Snake Weed test": fmt_ms(e, "recall", 7),
               "độ trễ p95 b1 (ms)": lat.get(FINAL_LATENCY[exp], {}).get("p95", np.nan)}
        if exp != "T00" and base:
            d = e["macro_f1"]["mean"] - base["macro_f1"]["mean"]
            s = max(e["macro_f1"]["std"], base["macro_f1"]["std"])
            row["Δ macro-F1 test so với mốc"] = f"{d:+.4f} (s = {s:.4f}, Δ/s = {d / s:.1f})"
        rows.append(row)
    return pd.DataFrame(rows)


def sheet_decomposition():
    """Phân tách đóng góp công thức huấn luyện × suy luận trên VAL (code/diag_factorial_val.py), 3 seed mỗi ô."""
    p = Path("logs/diag_factorial_val.json")
    if not p.exists():
        return pd.DataFrame()
    d = json.loads(p.read_text())
    names = {"I00_crop224": "I00: Resize256 + CenterCrop224", "I04_full288": "I04_288: ảnh đầy đủ resize 288"}
    rows = []
    for key, v in d["summary_macro_f1"].items():
        exp, inf = key.split("|")
        rows.append({"hạng mục": "ô", "công thức": exp, "suy luận": names[inf],
                     **{f"macro-F1 val seed {s}": x for s, x in enumerate(v["values"])},
                     "mean": v["mean"], "std": v["std"]})
    for exp, v in d["effect_inference"].items():
        rows.append({"hạng mục": "hiệu ứng suy luận (288 - 224, cùng checkpoint)", "công thức": exp,
                     "suy luận": "I04_288 - I00",
                     **{f"macro-F1 val seed {s}": x for s, x in enumerate(v["values"])},
                     "mean": v["mean"], "std": v["std"]})
    for inf, v in d["effect_recipe"].items():
        rows.append({"hạng mục": "hiệu ứng công thức (F01 - T00, hiệu hai trung bình; std = std lớn hơn của hai ô)",
                     "công thức": "F01 - T00",
                     "suy luận": names[inf], "mean": v["delta"], "std": v["s"]})
    rows.append({"hạng mục": "tương tác (hiệu ứng công thức ở 288 - ở 224)", "công thức": "", "suy luận": "",
                 "mean": d["interaction"]})
    return pd.DataFrame(rows)


def sheet_perclass():
    frames = []
    for exp, label in (("F01", "chung kết F01"), ("T00", "mốc T00+I00")):
        p = Path(f"eval_out/{exp}_per_class.csv")
        if not p.exists():
            continue
        d = pd.read_csv(p)
        frames.append(pd.DataFrame({
            "cấu hình": label, "lớp": d["class"], "số ảnh test": d["support"],
            "precision (mean)": d["precision_mean"], "precision (std)": d["precision_std"],
            "recall (mean)": d["recall_mean"], "recall (std)": d["recall_std"],
            "F1 (mean)": d["f1_mean"], "F1 (std)": d["f1_std"]}))
    return pd.concat(frames) if frames else pd.DataFrame()


def read_json(path: str) -> dict | None:
    p = Path(path)
    return json.loads(p.read_text()) if p.exists() else None


def ms(v: list) -> str:
    return f"{np.mean(v):.4f} ± {np.std(v, ddof=1):.4f}"


def sheet_amp_dtype():
    b = read_json("logs/bench_amp_dtype.json")
    if not b:
        return pd.DataFrame()
    return pd.DataFrame([{"mô hình": b["model"], "thiết bị": b["device"], "dtype": k, "batch": b["batch"],
                          "độ phân giải": b["img_size"], "tốc độ train (ảnh/s)": v["train_img_per_s"],
                          "thời gian một bước (ms)": v["step_ms"], "torch": b["torch"]} for k, v in b["results"].items()])


def sheet_bonus_folds():
    """Điểm thưởng: chung kết F01 (seed 0) trên fold 0, 1, 2; số test tính lại từ predictions/ bằng eval.py."""
    rows = []
    for k in (0, 1, 2):
        exp = "F01" if k == 0 else f"F01_fold{k}"
        st = read_json(f"logs/step4_{exp}.json")
        pred = Path(f"predictions/{exp}_seed0_test.csv")
        if not st or "0" not in st or not pred.exists():
            continue
        rec = st["0"]
        m = test_metrics(pred)
        rows.append({"exp_id": exp, "fold": k, "seed": 0, "số ảnh test": m["n"],
                     "T (khớp trên val của fold)": rec["temperature"], "macro-F1 val": rec["val_macro_f1"],
                     "macro-F1 test": m["macro_f1"], "top-1 test": m["top1"], "ECE test": m["ece"],
                     "recall Chinee Apple test": m["recall"][0], "recall Snake Weed test": m["recall"][7]})
    if len(rows) >= 2:
        df = pd.DataFrame(rows)
        agg = {"exp_id": f"mean ± std qua {len(rows)} fold", "fold": "tổng hợp", "seed": 0}
        for c in df.columns[4:]:
            agg[c] = ms(df[c].tolist())
        rows.append(agg)
    return pd.DataFrame(rows)


def sheet_bonus_probe(bb, tr):
    d = read_json("logs/bonus_dinov2_probe.json")
    if not d:
        return pd.DataFrame()
    label = {"dinov2_vits14": "DINOv2 ViT-S/14 (tự giám sát LVD-142M), đóng băng + hồi quy logistic",
             "deit_s_in1k": "DeiT-S (có giám sát ImageNet-1k), đóng băng + hồi quy logistic (cùng giao thức)"}
    rows = [{"mô hình": label[k], "timm": v["timm"], "cách dùng": "linear probe (đặc trưng đóng băng)",
             "#tham số (M)": v["params_m"], "C (CV 5-fold trên train)": v["best_C"], "macro-F1 val": v["val_macro_f1"],
             "top-1 val": v["val_top1"], "ECE val": v["val_ece"], "F1 Chinee Apple val": v["val_f1_chinee"],
             "F1 Snake Weed val": v["val_f1_snake"]} for k, v in d["results"].items()]
    ref = [("T01", tr, "DeiT-S đóng băng, head train 10 epoch có augmentation (Bước 2)"),
           ("B03", bb, "DeiT-S tinh chỉnh toàn bộ (Bước 1)"), ("B02", bb, "ConvNeXt-T tinh chỉnh toàn bộ (Bước 1)"),
           ("B01", bb, "ResNet-50 tinh chỉnh toàn bộ (Bước 1)"), ("B06", bb, "MobileNetV3 tinh chỉnh toàn bộ (Bước 1)")]
    for exp, df, desc in ref:
        r = df[df["exp_id"] == exp].iloc[0]
        rows.append({"mô hình": f"{exp}: {desc}", "cách dùng": "tham chiếu", "macro-F1 val": r["macro-F1 val"],
                     "top-1 val": r["top-1 val"], "F1 Chinee Apple val": r["F1 Chinee Apple val"],
                     "F1 Snake Weed val": r["F1 Snake Weed val"]})
    return pd.DataFrame(rows)


SHIFT_NAMES = {"clean": "sạch", "dark_0.5": "thiếu sáng ×0,5", "dark_0.3": "thiếu sáng ×0,3", "blur_1.5": "mờ σ = 1,5 px",
               "blur_3": "mờ σ = 3 px", "noise_0.05": "nhiễu σ = 0,05", "noise_0.10": "nhiễu σ = 0,10"}


def sheet_bonus_shift():
    d = read_json("logs/bonus_shift.json")
    if not d:
        return pd.DataFrame()
    rows = []
    for exp, res in d["results"].items():
        for cond, r in res.items():
            s = r["summary"]
            rows.append({"cấu hình": exp, "suy luận": f"{d['groups'][exp]['eval']} {d['groups'][exp]['img_size']}",
                         "điều kiện (val)": SHIFT_NAMES[cond], "số seed": len(r["per_seed"]),
                         "macro-F1 val": s["macro_f1"]["mean"], "std macro-F1": s["macro_f1"]["std"],
                         "top-1 val": s["top1"]["mean"], "F1 Chinee Apple": s["f1_chinee"]["mean"],
                         "F1 Snake Weed": s["f1_snake"]["mean"], "ECE (T = 1)": s["ece_T1"]["mean"],
                         "ECE (T khớp val sạch)": s["ece_Tclean"]["mean"], "ECE (T oracle, khớp lại trên ảnh lệch)":
                         s["ece_Toracle"]["mean"], "T oracle": s["T_oracle"]["mean"],
                         "độ tin cậy TB (T khớp val sạch)": s["conf_Tclean"]["mean"]})
    return pd.DataFrame(rows)


def sheet_bonus_tta():
    d = read_json("logs/bonus_tta.json")
    if not d:
        return pd.DataFrame()
    rows = []
    for exp, res in d.get("bn", {}).get("results", {}).items():
        for cond, r in res.items():
            rows.append({"phương pháp": "BN-adapt: thống kê BN tính lại trên chính ảnh lệch (không nhãn)",
                         "mô hình": exp, "điều kiện (val)": SHIFT_NAMES[cond],
                         "macro-F1 trước": r["goc"]["macro_f1"], "macro-F1 precise-BN (train sạch)":
                         r["precise_bn"]["macro_f1"], "macro-F1 sau": r["bn_adapt"]["macro_f1"],
                         "ECE trước": r["goc"]["ece"], "ECE sau": r["bn_adapt"]["ece"]})
    for cond, r in d.get("tent", {}).get("results", {}).items():
        rows.append({"phương pháp": "Tent: cập nhật γ/β LayerNorm, cực tiểu entropy, online 1 lượt (không nhãn)",
                     "mô hình": "F01 seed 0", "điều kiện (val)": SHIFT_NAMES[cond],
                     "macro-F1 trước": r["truoc"]["macro_f1"], "macro-F1 sau": r["tent"]["macro_f1"],
                     "ECE trước": r["truoc_Tclean"]["ece"], "ECE sau": r["tent_Tclean"]["ece"],
                     "ghi chú": "ECE với T khớp trên val sạch"})
    return pd.DataFrame(rows)


def sheet_bonus_onnx():
    d = read_json("logs/bonus_onnx.json")
    if not d:
        return pd.DataFrame()
    names = {"pytorch_mps": "PyTorch, GPU (MPS)", "pytorch_cpu": "PyTorch, CPU", "onnxruntime_cpu": "ONNX Runtime, CPU",
             "onnxruntime_coreml": "ONNX Runtime, CoreML EP"}
    rows = []
    for k, lat in d["latency_b1"].items():
        par = d["parity_val"].get(k, {})
        rows.append({"backend": names[k], "dtype": "fp32", "batch": 1, "độ phân giải": 288, "p50 (ms)": lat["p50"],
                     "p95 (ms)": lat["p95"], "p99 (ms)": lat["p99"], "số lần đo": lat["n"], "warmup": lat["warmup"],
                     "macro-F1 val": par.get("macro_f1", np.nan),
                     "sai khác logit lớn nhất so với PyTorch CPU": par.get("max_abs_logit_diff_vs_pytorch_cpu", np.nan),
                     "tỉ lệ trùng nhãn đoán": par.get("argmax_agreement_vs_pytorch_cpu", np.nan)})
    return pd.DataFrame(rows)


def sheet_bonus_attention():
    d = read_json("logs/bonus_attention.json")
    if not d:
        return pd.DataFrame()
    rows = []
    for key, label in (("sai", "đoán sai"), ("dung", "đoán đúng")):
        s = d["val_quantitative"][key]
        rows.append({"nhóm ảnh val (Chinee apple + Snake weed)": label, "số ảnh": s["n"],
                     "diện tích thực vật xanh (ExG-Otsu)": s["veg_area"]["mean"],
                     "Grad-CAM: tập trung vào thực vật (tỉ số)": s["gradcam_veg_ratio"]["mean"],
                     "Grad-CAM: entropy chuẩn hoá": s["gradcam_entropy"]["mean"],
                     "Grad-CAM: std entropy": s["gradcam_entropy"]["std"],
                     "rollout: tập trung vào thực vật (tỉ số)": s["rollout_veg_ratio"]["mean"],
                     "rollout: entropy chuẩn hoá": s["rollout_entropy"]["mean"]})
    return pd.DataFrame(rows)


def sheet_summary(bb, tr, inf):
    rows = []
    for _, r in bb.iterrows():
        rows.append({"exp_id": r["exp_id"], "loại": "backbone", "mô tả": f"{r['backbone']} (công thức nền)",
                     "macro-F1 val": r["macro-F1 val"], "top-1 val": r["top-1 val"],
                     "độ trễ p95 b1 (ms)": r["độ trễ batch-1 p95 (ms)"], "GMAC": r["GMAC"],
                     "train/epoch (s)": r["thời gian train/epoch (s)"]})
    for _, r in tr[tr["so với nhiễu"] != "nhiễu seed (cùng cấu hình)"].iterrows():  # bỏ các dòng lặp seed của T00
        rows.append({"exp_id": r["exp_id"], "loại": "huấn luyện", "mô tả": r["khác T00 ở điểm nào"],
                     "macro-F1 val": r["macro-F1 val"], "top-1 val": r["top-1 val"],
                     "train/epoch (s)": r["thời gian train/epoch (s)"]})
    for _, r in inf.iterrows():
        rows.append({"exp_id": r["exp_id"], "loại": "suy luận", "mô tả": r["phương pháp"],
                     "macro-F1 val": r["macro-F1 val"], "top-1 val": r["top-1 val"],
                     "độ trễ p95 b1 (ms)": r["độ trễ p95 b1 (ms)"]})
    df = pd.DataFrame(rows).sort_values("macro-F1 val", ascending=False).head(10).reset_index(drop=True)
    df.insert(0, "hạng", np.arange(1, len(df) + 1))
    df["ghi chú"] = df["exp_id"].map(SUMMARY_NOTES).fillna("")
    return df


def write_table(ws, df: pd.DataFrame, hdr: int, highlight: int | None):
    """Định dạng một bảng đã ghi ở hàng tiêu đề `hdr` (1-based): tiêu đề đậm, số 4 chữ số, tô dòng `highlight`."""
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    for c in ws[hdr]:
        c.font = Font(bold=True)
    for j, col in enumerate(df.columns, 1):
        width = max([len(str(col))] + [len(f"{v:.4f}" if isinstance(v, float) else str(v)) for v in df[col].head(50)])
        letter = get_column_letter(j)
        ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0, min(60, width + 2))
        for i in range(len(df)):
            cell = ws.cell(row=hdr + 1 + i, column=j)
            if isinstance(cell.value, float):
                cell.number_format = "0.0000"
    if highlight is not None:
        for j in range(1, len(df.columns) + 1):
            ws.cell(row=hdr + 1 + highlight, column=j).fill = PatternFill("solid", fgColor="FFF2CC")


def best_row(name: str, df: pd.DataFrame) -> int | None:
    """Dòng tô vàng: macro-F1 val cao nhất (sheet Training: chỉ xét các ablation, bỏ mốc và các dòng đo nhiễu)."""
    if "macro-F1 val" not in df.columns or not len(df):
        return None
    vals = pd.to_numeric(df["macro-F1 val"], errors="coerce")
    if name == "Training":
        vals = vals.where(df["exp_id"] != "T00")
    return int(vals.idxmax()) if vals.notna().any() else None


def write_xlsx(sheets: dict, path: str, notes: list[str], compare: pd.DataFrame):
    from openpyxl.styles import Font

    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            start = 0
            if name == "Summary":
                # ghi chú, rồi bảng chung kết so với mốc (TEST), rồi top 10 theo macro-F1 val
                compare.to_excel(xw, sheet_name=name, index=False, startrow=len(notes) + 1)
                start = len(notes) + len(compare) + 4
            df.to_excel(xw, sheet_name=name, index=False, startrow=start)
            ws = xw.sheets[name]
            if name == "Summary":
                for i, line in enumerate(notes):
                    ws.cell(row=i + 1, column=1, value=line).font = Font(bold=(i == 0))
                ws.cell(row=len(notes) + 1, column=1, value="Chung kết so với mốc trên TEST (test chạy một lần mỗi "
                        "seed; số theo eval.py)").font = Font(bold=True)
                write_table(ws, compare, len(notes) + 2, 0 if len(compare) else None)
                ws.cell(row=start, column=1, value="Top 10 cấu hình theo macro-F1 VAL (Bước 1-3, 1 seed)").font = \
                    Font(bold=True)
            hdr = start + 1
            ws.freeze_panes = ws.cell(row=hdr + 1, column=1) if name != "Summary" else None
            write_table(ws, df, hdr, best_row(name, df))


def main():
    bb = sheet_backbones()
    tr = sheet_training()
    inf, latency = sheet_inference()
    fin = sheet_final()
    pc = sheet_perclass()
    dec = sheet_decomposition()
    summ = sheet_summary(bb, tr, inf)
    compare = sheet_compare()
    notes = ["Summary - chung kết so với mốc (TEST) và top 10 cấu hình theo macro-F1 VAL; chi tiết ở các sheet khác"]
    grade = Path("eval_out/grade_I.json")
    for exp in ("F01", "T00"):
        p = Path(f"eval_out/{exp}_summary.json")
        if p.exists():
            e = json.loads(p.read_text())
            notes.append(f"{exp} TEST ({len(e['seeds'])} seed): macro-F1 {e['macro_f1']['mean']:.4f} ± "
                         f"{e['macro_f1']['std']:.4f} | top-1 {e['top1']['mean']:.4f} ± {e['top1']['std']:.4f} | "
                         f"ECE {e['ece']['mean']:.4f} ± {e['ece']['std']:.4f}  ({FINAL_META.get(exp, '')})")
    if grade.exists():
        g = json.loads(grade.read_text())
        notes.append("eval.py grade (phần I, đề xuất): " + "; ".join(
            f"{i['code']}={i['points']}/{i['max']}" for i in g["items"]) + f" -> {g['total']}/{g['max_scored']}")
    notes.append("Dòng tô vàng: cấu hình chung kết (bảng đầu) và macro-F1 val cao nhất (các bảng khác; sheet "
                 "Training chỉ xét ablation). Đơn vị: độ trễ ms, thời gian s; thiết bị Apple M5 GPU (MPS), torch 2.14.1.")
    sd = noise_std()
    notes.append(f"Nhiễu seed: std macro-F1 val của T00 qua 3 seed σ = {sd:.4f}; hiệu hai lần chạy 1 seed có std "
                 f"≈ √2·σ = {np.sqrt(2) * sd:.4f}, nên |Δ| nhỏ hơn mức này là không phân biệt được.")
    sheets = {"Summary": summ, "Backbones": bb, "Training": tr, "Inference": inf, "Final": fin,
              "PerClass": pc, "Latency": latency}
    if len(dec):
        sheets["Decomposition"] = dec
    # điểm thưởng (RUBRIC mục 2) và căn cứ chọn dtype AMP; sheet nào chưa có log thì bỏ qua
    extra = {"AMP_dtype": sheet_amp_dtype(), "Bonus_Folds": sheet_bonus_folds(), "Bonus_Probe": sheet_bonus_probe(bb, tr),
             "Bonus_Shift": sheet_bonus_shift(), "Bonus_TTA": sheet_bonus_tta(), "Bonus_Attention": sheet_bonus_attention(),
             "Bonus_ONNX": sheet_bonus_onnx()}
    sheets.update({k: v for k, v in extra.items() if len(v)})
    write_xlsx(sheets, "results.xlsx", notes, compare)
    print("saved results.xlsx", {k: len(v) for k, v in sheets.items()})


if __name__ == "__main__":
    main()
